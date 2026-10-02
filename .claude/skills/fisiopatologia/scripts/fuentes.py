"""Consulta de fuentes abiertas para el material de fisiopatología de una enfermedad.

Cada función devuelve datos verificables y registra la procedencia en
`assets/ilustraciones/registro.json` cuando guarda un archivo.

Uso rápido desde la terminal:
  python3 fuentes.py pubchem darolutamida
  python3 fuentes.py chembl darolutamide
  python3 fuentes.py pdb-buscar "androgen receptor darolutamide"
  python3 fuentes.py pdb-ligandos 2AMA
  python3 fuentes.py bioicons testis
  python3 fuentes.py servier-kits
  python3 fuentes.py servier-diapositivas Endocrinology
  python3 fuentes.py servier-extraer Endocrinology 11 "Group 641" testiculo
  python3 fuentes.py cima atorvastatina          # ficha técnica en español (AEMPS)
  python3 fuentes.py dailymed atorvastatin       # ficha técnica de la FDA
  python3 fuentes.py uniprot SREBF2
  python3 fuentes.py reactome "SREBP cholesterol"
  python3 fuentes.py europepmc "SREBP-2 AND LDL receptor AND statin" "SREBP-?2.*(LDLR|LDL receptor)"
  python3 fuentes.py pdf <URL del PDF> "<regex>"   # frases de un PDF público con su página
  python3 fuentes.py guias "sepsis" [--titulo] [--latam]   # guías recientes, también las aún sin indexar
  python3 fuentes.py vigencia 34599691              # ¿hay una versión posterior de esta guía? (y erratas)
  python3 fuentes.py texto 27418577 "MRSA"          # texto completo: PMC y, si no, acceso abierto (Unpaywall)
  python3 fuentes.py iris "AWaRe antibiotic book" [--ops]   # documentos de la OMS (o de la OPS) con su PDF
  python3 fuentes.py binasss "infecciones"          # protocolos y normas de la CCSS (BINASSS, Costa Rica)
  python3 fuentes.py ema jardiance                  # medicamento de la EMA: estado, fechas y página del EPAR
  python3 fuentes.py ensayos empagliflozin          # ensayos aleatorizados (fase III primero) con su NCT
  python3 fuentes.py ensayo NCT01131676             # registro del ensayo en ClinicalTrials.gov
  python3 fuentes.py --registro ejemplos/<tema> pubmed "..."   # anota la consulta en busquedas.jsonl
  python3 fuentes.py uniprot-proteina "beta-lactamase" 1280   # proteína de S. aureus (taxón 1280)

Si una fuente no responde (red bloqueada), las funciones lanzan RuntimeError con el
nombre del dominio para avisar a la persona usuaria; no inventan datos.
"""
import html
import json
import os
import re
import subprocess
import sys
import urllib.parse
import urllib.request
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
ILUSTRACIONES = RAIZ / "assets" / "ilustraciones"
REGISTRO = ILUSTRACIONES / "registro.json"
CACHE = Path("/tmp/fisiopatologia-cache")
CACHE.mkdir(exist_ok=True)
AGENTE = "fisiopatologia/1.0"  # User-Agent de las consultas

LICENCIAS = {
    "servier": "CC BY 4.0 (Servier Medical Art, smart.servier.com)",
    "servier-bioicons": "CC BY 3.0 (Servier Medical Art vía Bioicons)",
    "pubchem": "Dominio público (NCBI PubChem)",
    "chembl": "CC BY-SA 3.0 (EMBL-EBI ChEMBL)",
    "pdb": "CC0 (RCSB Protein Data Bank)",
    "alphafold": "CC BY 4.0 (AlphaFold DB, EMBL-EBI/Google DeepMind)",
    "nih-bioart": "Dominio público salvo indicación (NIH BioArt, bioart.niaid.nih.gov)",
    "bioicons": "Según carpeta de licencia del icono (Bioicons)",
    "togotv": "CC BY 4.0 (© DBCLS TogoTV, togotv.dbcls.jp)",
    "commons": "Según la ficha del archivo en Wikimedia Commons",
}


_NAVEGADOR = {"User-Agent": "Mozilla/5.0 (compatible; skills-docencia/1.0)", "Accept-Language": "es,en;q=0.8,*;q=0.5"}

# --- Registro de búsquedas -------------------------------------------------------------------------------------
# Con registrar_busquedas(carpeta) o la opción --registro <carpeta>, cada consulta queda anotada (fecha, función,
# argumentos y qué devolvió) en <carpeta>/busquedas.jsonl. `bibliografia.py busqueda` lo resume en «Cómo se buscó».
_REGISTRO_BUSQUEDAS = os.environ.get("FUENTES_REGISTRO") or None
_PROFUNDIDAD = [0]  # solo se anotan las consultas que pide la skill, no las que hacen otras funciones por dentro


def registrar_busquedas(destino):
    """Empieza a anotar las consultas en `destino` (archivo .jsonl o carpeta, donde crea busquedas.jsonl)."""
    global _REGISTRO_BUSQUEDAS
    ruta = Path(destino)
    if ruta.is_dir() or not ruta.suffix:
        ruta.mkdir(parents=True, exist_ok=True)
        ruta = ruta / "busquedas.jsonl"
    _REGISTRO_BUSQUEDAS = str(ruta)
    return ruta


def _resumen_resultado(resultado):
    """Lo que se anota de cada resultado: cuántos hubo y sus identificadores, no el contenido."""
    if isinstance(resultado, list):
        ids = [x.get("pmid") or x.get("url") or x.get("id") or x.get("uniprot") or x.get("reactome")
               for x in resultado[:25] if isinstance(x, dict)]
        return {"n": len(resultado), "ids": [i for i in ids if i]}
    if isinstance(resultado, dict):
        claves = ("pmid", "pmcid", "doi", "url", "fuente", "texto_completo", "paginas", "set_id", "nregistro",
                  "es_acceso_abierto")
        salida = {k: resultado[k] for k in claves if resultado.get(k) is not None}
        if isinstance(resultado.get("fragmentos"), dict):
            salida["fragmentos"] = {p: len(v) for p, v in resultado["fragmentos"].items()}
        if isinstance(resultado.get("posteriores"), list):
            salida["posteriores"] = [c.get("pmid") for c in resultado["posteriores"]]
        return salida
    return {}


def _registrada(funcion):
    """Anota la consulta en el registro de búsquedas, si está activo."""
    import functools

    @functools.wraps(funcion)
    def envoltura(*args, **kwargs):
        _PROFUNDIDAD[0] += 1
        try:
            resultado = funcion(*args, **kwargs)
        finally:
            _PROFUNDIDAD[0] -= 1
        if _REGISTRO_BUSQUEDAS and _PROFUNDIDAD[0] == 0:
            from datetime import datetime, timezone
            linea = {"fecha": datetime.now(timezone.utc).isoformat(timespec="seconds"), "funcion": funcion.__name__,
                     "argumentos": [str(a) for a in args], "opciones": {k: str(v) for k, v in kwargs.items()},
                     "resultado": _resumen_resultado(resultado)}
            with open(_REGISTRO_BUSQUEDAS, "a", encoding="utf-8") as archivo:
                archivo.write(json.dumps(linea, ensure_ascii=False) + "\n")
        return resultado
    return envoltura


def _cortesia_ncbi(url):
    """Añade a las consultas a NCBI lo que pide su política de uso: `tool` y, si están en el entorno, `email`
    (NCBI_EMAIL) y `api_key` (NCBI_API_KEY; con clave, NCBI admite 10 consultas por segundo en lugar de 3)."""
    partes = urllib.parse.urlsplit(url)
    if partes.netloc not in ("eutils.ncbi.nlm.nih.gov", "pmc.ncbi.nlm.nih.gov") or "tool=" in partes.query:
        return url
    extra = {"tool": "skills-docencia"}
    for variable, parametro in (("NCBI_EMAIL", "email"), ("NCBI_API_KEY", "api_key")):
        if os.environ.get(variable):
            extra[parametro] = os.environ[variable]
    consulta = "&".join(x for x in (partes.query, urllib.parse.urlencode(extra)) if x)
    return urllib.parse.urlunsplit(partes._replace(query=consulta))


def _get(url, timeout=60, datos=None, cabeceras=None, intentos=4, con_url=False):
    """Descarga una URL. Reintenta con espera creciente si el servidor limita la frecuencia (429), falla (5xx)
    o se corta la conexión; ante un 429 respeta la cabecera Retry-After (hasta 30 s). Con con_url=True devuelve
    también la dirección final, tras las redirecciones."""
    import time
    import urllib.error
    url = _cortesia_ncbi(url)
    req = urllib.request.Request(url, data=datos, headers=cabeceras or {"User-Agent": AGENTE})
    for intento in range(intentos):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                contenido = r.read()
                return (contenido, r.geturl()) if con_url else contenido
        except urllib.error.HTTPError as e:
            transitorio = e.code == 429 or e.code >= 500 or (e.code == 400 and "eutils.ncbi" in url)  # NCBI
            if transitorio and intento < intentos - 1:
                espera = e.headers.get("Retry-After", "") if e.headers else ""
                time.sleep(min(int(espera), 30) if espera.isdigit() else 2 ** intento)
                continue
            error = e
        except (urllib.error.URLError, ConnectionError, TimeoutError) as e:
            if intento < intentos - 1:  # corte de red o túnel cerrado: suele ser pasajero
                time.sleep(2 ** intento)
                continue
            error = e
        except Exception as e:  # noqa: BLE001
            error = e
        break
    dominio = urllib.parse.urlparse(url).netloc
    raise RuntimeError(f"No se pudo consultar {dominio}: {error}. Puede que la red del entorno lo bloquee.") from error


def registrar(archivo, fuente, origen, notas=""):
    """Anota procedencia y licencia de un archivo incorporado a la biblioteca."""
    registro = json.loads(REGISTRO.read_text()) if REGISTRO.exists() else {}
    registro[Path(archivo).name] = {"fuente": fuente, "licencia": LICENCIAS.get(fuente, fuente), "origen": origen,
                                   "fecha": date.today().isoformat(), "notas": notas}
    REGISTRO.write_text(json.dumps(registro, indent=2, ensure_ascii=False))


# --- Química -------------------------------------------------------------------

@_registrada
def pubchem(nombre):
    """SMILES, fórmula y CID de un compuesto por nombre (inglés o DCI)."""
    q = urllib.parse.quote(nombre)
    url = (f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/{q}/property/"
           "IsomericSMILES,CanonicalSMILES,MolecularFormula,IUPACName/JSON")
    p = json.loads(_get(url))["PropertyTable"]["Properties"][0]
    smiles = p.get("IsomericSMILES") or p.get("SMILES") or p.get("CanonicalSMILES") or p.get("ConnectivitySMILES")
    return {"cid": p["CID"], "smiles": smiles, "formula": p["MolecularFormula"], "iupac": p.get("IUPACName"),
            "fuente": f"PubChem CID {p['CID']}"}


@_registrada
def chembl(nombre):
    """Mecanismo de acción y diana documentados en ChEMBL.

    Busca primero por nombre exacto y, si no, por texto; usa la molécula «padre» porque
    las sales (p. ej., atorvastatina cálcica) no tienen mecanismos asociados.
    """
    q = urllib.parse.quote(nombre)
    base = "https://www.ebi.ac.uk/chembl/api/data"
    mols = json.loads(_get(f"{base}/molecule.json?pref_name__iexact={q}&limit=1"))["molecules"]
    mols += json.loads(_get(f"{base}/molecule/search.json?q={q}&limit=5"))["molecules"]
    ids = []
    for m in mols:
        padre = (m.get("molecule_hierarchy") or {}).get("parent_chembl_id") or m["molecule_chembl_id"]
        if padre not in ids:
            ids.append(padre)
    for chembl_id in ids:
        # El mecanismo puede estar registrado en una sal (molécula hija); se busca por la molécula padre.
        mecs = json.loads(_get(f"{base}/mechanism.json?parent_molecule_chembl_id={chembl_id}"))["mechanisms"]
        if not mecs:
            continue
        salida = []
        for m in mecs:
            diana = json.loads(_get(f"{base}/target/{m['target_chembl_id']}.json"))
            salida.append({"chembl_id": chembl_id, "mecanismo": m["mechanism_of_action"], "accion": m["action_type"],
                           "diana": diana["pref_name"], "tipo_diana": diana["target_type"],
                           "referencias": [r.get("ref_url") for r in m.get("mechanism_refs", [])]})
        return salida
    return []


# --- Estructuras de proteínas -------------------------------------------------

@_registrada
def pdb_buscar(texto, filas=10):
    consulta = {"query": {"type": "terminal", "service": "full_text", "parameters": {"value": texto}},
                "return_type": "entry", "request_options": {"paginate": {"start": 0, "rows": filas}}}
    r = json.loads(_get("https://search.rcsb.org/rcsbsearch/v2/query", datos=json.dumps(consulta).encode(),
                        cabeceras={"Content-Type": "application/json"}))
    salida = []
    for e in r.get("result_set", []):
        d = json.loads(_get(f"https://data.rcsb.org/rest/v1/core/entry/{e['identifier']}"))
        salida.append({"pdb": e["identifier"], "titulo": d["struct"]["title"]})
    return salida


@_registrada
def pdb_ligandos(pdb_id):
    d = json.loads(_get(f"https://data.rcsb.org/rest/v1/core/entry/{pdb_id}"))
    ids = d["rcsb_entry_container_identifiers"].get("non_polymer_entity_ids") or []
    salida = []
    for e in ids:
        n = json.loads(_get(f"https://data.rcsb.org/rest/v1/core/nonpolymer_entity/{pdb_id}/{e}"))["pdbx_entity_nonpoly"]
        salida.append({"codigo": n["comp_id"], "nombre": n["name"]})
    return salida


def pdb_descargar(pdb_id):
    destino = CACHE / f"{pdb_id.upper()}.pdb"
    if not destino.exists():
        destino.write_bytes(_get(f"https://files.rcsb.org/download/{pdb_id.upper()}.pdb", timeout=120))
    return destino


def alphafold_descargar(uniprot):
    destino = CACHE / f"AF-{uniprot}.pdb"
    if not destino.exists():
        info = json.loads(_get(f"https://alphafold.ebi.ac.uk/api/prediction/{uniprot}"))[0]
        destino.write_bytes(_get(info["pdbUrl"], timeout=120))
    return destino


# --- Ilustraciones -------------------------------------------------------------

@_registrada
def bioicons(termino, repo=CACHE / "bioicons"):
    """Busca iconos en Bioicons (GitHub). Devuelve rutas con su carpeta de licencia."""
    if not repo.exists():
        subprocess.run(["git", "clone", "-q", "--depth", "1", "--filter=blob:none", "--sparse",
                        "https://github.com/duerrsimon/bioicons.git", str(repo)], check=True, timeout=600)
        subprocess.run(["git", "-C", str(repo), "sparse-checkout", "set", "static/icons"], check=True, timeout=600)
    base = repo / "static" / "icons"
    return [str(p.relative_to(base)) for p in base.rglob("*.svg") if termino.lower() in p.name.lower()]


@_registrada
def togopic(termino, maximo=30):
    """Busca en la galería de TogoTV (DBCLS, Japón; CC BY 4.0). La búsqueda mira el nombre y las etiquetas:
    en inglés solo encuentra por el nombre; en japonés también por la etiqueta (感染症, 細菌, ウイルス, 真菌…)."""
    q = urllib.parse.urlencode({"target": "pictures", "text": termino})
    datos = json.loads(_get(f"https://togotv-api.dbcls.jp/api/search?{q}"))["data"][:maximo]
    return [{"nombre": d.get("name_en"), "nombre_ja": d.get("name"), "svg": d.get("svg"), "doi": d.get("id"),
             "licencia": d.get("license"), "autoria": d.get("author_str"), "fecha": d.get("uploadDate"),
             "etiquetas": d.get("other_tags_comma_en")} for d in datos if d.get("svg")]


def togopic_descargar(svg, nombre, doi=""):
    """Descarga un SVG de TogoTV a assets/ilustraciones/togotv-<nombre>.svg y lo registra.
    svg: el campo «svg» de togopic(); doi: el campo «doi» (identifica la imagen en el registro)."""
    destino = ILUSTRACIONES / f"togotv-{nombre}.svg"
    destino.write_bytes(_get("https://dbarchive.biosciencedbc.jp/data/togo-pic/image/" + urllib.parse.quote(svg),
                             timeout=120))
    registrar(destino, "togotv", f"TogoTV {svg} {doi}".strip())
    return destino


def _commons_info(paginas):
    salida = []
    for pagina_ in paginas:
        info = pagina_["imageinfo"][0]
        meta = info.get("extmetadata", {})
        licencia = meta.get("LicenseShortName", {}).get("value", "")
        if re.search(r"\bN[CD]\b|NonCommercial|NoDeriv", licencia, re.I):
            continue
        salida.append({"titulo": pagina_["title"], "url": info["url"].split("?")[0], "licencia": licencia,
                       "autoria": _texto_plano(meta.get("Artist", {}).get("value", "")).strip()[:120],
                       "aviso": "CC BY-SA: avisa antes de usarla" if "SA" in licencia else None})
    return salida


def _commons_api(parametros):
    q = urllib.parse.urlencode({"action": "query", "format": "json", "prop": "imageinfo",
                                "iiprop": "url|extmetadata", **parametros})
    datos = json.loads(_get(f"https://commons.wikimedia.org/w/api.php?{q}", cabeceras=COMMONS_UA))
    return list(datos.get("query", {}).get("pages", {}).values())


COMMONS_UA = {"User-Agent": "skills-docencia/1.0 (https://github.com/davidtencio/skills; material educativo)"}


@_registrada
def commons(termino, maximo=20):
    """Busca dibujos y diagramas en Wikimedia Commons, en cualquier idioma (p. ej., «Pharmakokinetik»,
    «pharmacocinétique»). Devuelve la licencia y la autoría de cada archivo; descarta las licencias NC y ND."""
    return _commons_info(_commons_api({"generator": "search", "gsrnamespace": 6,
                                       "gsrsearch": f"{termino} filetype:drawing", "gsrlimit": maximo}))


def commons_descargar(titulo, nombre):
    """Descarga un archivo de Wikimedia Commons («File:…») a assets/ilustraciones/commons-<nombre>.<ext>
    y lo registra con su licencia y autoría. Rechaza las licencias NC y ND."""
    archivos = _commons_info(_commons_api({"titles": titulo}))
    if not archivos:
        raise ValueError(f"{titulo}: no existe o su licencia no permite reutilizarlo (NC o ND)")
    r = archivos[0]
    destino = ILUSTRACIONES / f"commons-{nombre}{Path(r['url']).suffix}"
    destino.write_bytes(_get(r["url"], timeout=120, cabeceras=COMMONS_UA))
    registrar(destino, "commons", f"Wikimedia Commons, {r['titulo']}", f"{r['licencia']}; autoría: {r['autoria']}")
    return destino


def servier_kits():
    """Kits de PowerPoint de Servier Medical Art por categoría (vectoriales, CC BY 4.0)."""
    html = _get("https://smart.servier.com/image-kits-by-category/").decode("utf-8", "ignore")
    return sorted(set(re.findall(r"https://smart\.servier\.com/wp-content/uploads/[^\"]+\.pptx", html)))


def _kit(nombre):
    pptx = CACHE / f"SMART-{nombre}.pptx"
    if not pptx.exists():
        url = next(u for u in servier_kits() if u.endswith(f"SMART-{nombre}.pptx"))
        pptx.write_bytes(_get(url, timeout=300))
    return pptx


def servier_diapositivas(kit):
    """Lista diapositivas y grupos de dibujo de un kit, con sus títulos."""
    from pptx import Presentation
    p = Presentation(str(_kit(kit)))
    salida = []
    for i, s in enumerate(p.slides, 1):
        textos = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        grupos = [sh.name for sh in s.shapes if sh.shape_type == 6]
        if grupos:
            salida.append({"diapositiva": i, "textos": textos[:4], "grupos": grupos})
    return salida


def servier_extraer(kit, diapositiva, grupo, nombre, recorte=None):
    """Extrae un dibujo de un kit de Servier como SVG vectorial sin el fondo de la diapositiva.

    Requiere LibreOffice Impress (soffice), python-pptx y pymupdf.
    recorte: (x0, y0, x1, y1) en puntos de la diapositiva para quedarse con una parte del grupo.
    """
    import pymupdf
    from pptx import Presentation
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    p = Presentation(str(_kit(kit)))
    for m in p.slide_masters:
        for contenedor in [m] + list(m.slide_layouts):
            for sh in list(contenedor.shapes):
                sh._element.getparent().remove(sh._element)
            fondo = contenedor._element.find(f".//{ns}bg")
            if fondo is not None:
                fondo.getparent().remove(fondo)
    s = p.slides[diapositiva - 1]
    caja = None
    for sh in list(s.shapes):
        if sh.name == grupo:
            caja = (sh.left / 12700, sh.top / 12700, (sh.left + sh.width) / 12700, (sh.top + sh.height) / 12700)
        else:
            sh._element.getparent().remove(sh._element)
    if caja is None:
        raise ValueError(f"No existe el grupo {grupo!r} en la diapositiva {diapositiva}")
    tmp = CACHE / f"extraer-{nombre}.pptx"
    p.save(str(tmp))
    subprocess.run(["soffice", "--headless", "--norestore", "-env:UserInstallation=file:///tmp/lo_profile",
                    "--convert-to", "pdf", "--outdir", str(CACHE), str(tmp)], check=True, capture_output=True,
                   timeout=600)
    svg = pymupdf.open(str(tmp.with_suffix(".pdf")))[diapositiva - 1].get_svg_image(text_as_path=True)
    # Quita el rectángulo blanco de fondo que LibreOffice añade a cada página
    svg = re.sub(r'<path[^>]*d="M0 [\d.]+H[\d.]+V[\d.]+H0V[\d.]+Z" fill="#ffffff"[^>]*/>', "", svg)
    x0, y0, x1, y1 = recorte or caja
    svg = re.sub(r'viewBox="[^"]*"', f'viewBox="{x0 - 2:.1f} {y0 - 2:.1f} {x1 - x0 + 4:.1f} {y1 - y0 + 4:.1f}"', svg, 1)
    svg = re.sub(r'width="[^"]*" height="[^"]*"', f'width="{x1 - x0 + 4:.0f}" height="{y1 - y0 + 4:.0f}"', svg, 1)
    destino = ILUSTRACIONES / f"servier4-{nombre}.svg"
    destino.write_text(svg)
    registrar(destino, "servier", f"Kit SMART-{kit}.pptx, diapositiva {diapositiva}, {grupo}",
              f"recorte {recorte}" if recorte else "")
    return destino


# --- Verificación: fichas técnicas, biología curada y literatura ------------------

def _texto_plano(contenido):
    import html as _html
    contenido = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", contenido, flags=re.S)
    return _html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", contenido)))


@_registrada
def dailymed(nombre, secciones=("12.1 Mechanism of Action", "12.3 Pharmacokinetics", "12.4 Microbiology", "5.1 ",
                                 "7.1 "), largo=1200):
    """Ficha técnica de la FDA (DailyMed): devuelve fragmentos de las secciones pedidas.
    En los antimicrobianos, el mecanismo, la resistencia y la sensibilidad están en 12.4 Microbiology."""
    q = urllib.parse.quote(nombre)
    datos = json.loads(_get(f"https://dailymed.nlm.nih.gov/dailymed/services/v2/spls.json?drug_name={q}&pagesize=25"))
    if not datos["data"]:
        return {}
    # La del principio activo solo antes que una combinación («SYNJARDY (EMPAGLIFLOZIN AND METFORMIN …)»).
    def combinacion(d):  # los principios activos van entre los primeros paréntesis del título
        activos = re.search(r"\(([^)]*)\)", d["title"])
        return bool(activos and re.search(r"\bAND\b|/|,", activos.group(1)))
    spl = min(datos["data"], key=lambda d: (combinacion(d), -(d.get("spl_version") or 0)))
    texto = _texto_plano(_get(f"https://dailymed.nlm.nih.gov/dailymed/services/v2/spls/{spl['setid']}.xml",
                              timeout=120).decode("utf-8", "ignore"))
    salida = {"fuente": f"DailyMed, {spl['title'][:120]} (setid {spl['setid']})", "setid": spl["setid"],
              "version": spl.get("spl_version"), "fecha": spl.get("published_date"),
              "url": f"https://dailymed.nlm.nih.gov/dailymed/drugInfo.cfm?setid={spl['setid']}"}
    for sec in secciones:
        i = texto.find(sec)
        if i >= 0:
            salida[sec.strip()] = texto[i:i + largo]
    return salida


@_registrada
def cima(nombre, secciones=("4.1", "4.2", "4.5", "4.8", "5.1", "5.2", "5.3"), largo=1500):
    """Ficha técnica española (CIMA, AEMPS), en español, con la fecha del texto y las notas de seguridad de la AEMPS.
    Prefiere el medicamento original comercializado y sin combinar. nombre: principio activo o marca."""
    q = urllib.parse.quote(nombre)
    res = {}  # por principio activo (encuentra el original, p. ej. Jardiance) y por nombre del medicamento
    for campo in ("practiv1", "nombre"):
        for r in json.loads(_get(f"https://cima.aemps.es/cima/rest/medicamentos?{campo}={q}"))["resultados"]:
            res.setdefault(r["nregistro"], r)
    res = list(res.values())
    # el original comercializado antes que los genéricos y las combinaciones
    res = sorted(res, key=lambda r: ("/" in r["nombre"], bool(r.get("generico")), not r.get("comerc"), len(r["nombre"])))
    for r in res:
        ficha = next((d for d in r.get("docs", []) if d["tipo"] == 1 and d.get("urlHtml")), None)
        if not ficha:
            continue
        texto = _texto_plano(_get(ficha["urlHtml"], timeout=120).decode("utf-8", "ignore"))
        fecha = date.fromtimestamp(ficha["fecha"] / 1000).isoformat() if ficha.get("fecha") else None
        salida = {"fuente": f"CIMA (AEMPS), ficha técnica de {r['nombre']} (n.º registro {r['nregistro']})",
                  "url": ficha["urlHtml"], "nregistro": r["nregistro"], "fecha_ficha": fecha,
                  "titular": r.get("labtitular"), "comercializado": r.get("comerc"),
                  "autorizacion_europea": bool(r.get("ema"))}
        if r.get("notas"):  # notas de seguridad de la AEMPS sobre el medicamento
            try:
                salida["notas_seguridad"] = [
                    {"fecha": date.fromtimestamp(n["fecha"] / 1000).isoformat(), "referencia": n.get("referencia"),
                     "asunto": n.get("asunto"), "url": n.get("url")}
                    for n in json.loads(_get(f"https://cima.aemps.es/cima/rest/notas?nregistro={r['nregistro']}"))]
            except (RuntimeError, ValueError):
                salida["notas_seguridad"] = None
        for sec in secciones:
            m = re.search(rf"\b{re.escape(sec)}\.? [A-ZÁÉÍÓÚ][^0-9]{{3,60}}", texto)
            if m:
                salida[sec] = texto[m.start():m.start() + largo]
        return salida
    return {}


@_registrada
def uniprot(gen=None, organismo=9606, proteina=None, maximo=10):
    """Función de una proteína revisada en UniProt (con PMID de respaldo), por gen o por nombre de proteína.

    Por gen, se queda con la entrada cuyo nombre de gen principal coincide exactamente: algunos símbolos
    (p. ej., KLK3) también aparecen como sinónimos de otros genes. Por proteína (p. ej., «penicillin-binding
    protein 2a»), sirve también para bacterias, hongos y virus: organismo=None busca en todos, o pasa el
    identificador de taxonomía de NCBI (1280 = Staphylococcus aureus; 287 = Pseudomonas aeruginosa).
    """
    if not gen and not proteina:
        raise ValueError("Indica gen o proteina")
    partes = [f"gene:{gen}" if gen else f'protein_name:"{proteina}"', "reviewed:true"]
    if organismo:
        partes.append(f"taxonomy_id:{organismo}")  # incluye las cepas del taxón
    q = urllib.parse.quote(" AND ".join(partes))
    tsv = _get(f"https://rest.uniprot.org/uniprotkb/search?query={q}&fields=accession,gene_primary,protein_name,"
               f"organism_name,cc_function,cc_subcellular_location&format=tsv&size={maximo}").decode()
    filas = [f.split("\t") + [""] * 6 for f in tsv.strip().split("\n")[1:] if f]
    if gen:
        filas = [f for f in filas if f[1].upper() == gen.upper()] or filas
    return [{"uniprot": f[0], "gen": f[1], "proteina": f[2], "organismo": f[3], "funcion": f[4],
             "localizacion": f[5]} for f in filas]


@_registrada
def reactome(texto, especie="Homo sapiens"):
    """Vías de Reactome que coinciden con el texto, con su resumen."""
    q = urllib.parse.quote(texto)
    r = json.loads(_get(f"https://reactome.org/ContentService/search/query?query={q}&species="
                        f"{urllib.parse.quote(especie)}&types=Pathway"))
    salida = []
    for grupo in r.get("results", []):
        for e in grupo["entries"][:5]:
            d = json.loads(_get(f"https://reactome.org/ContentService/data/query/{e['stId']}"))
            resumen = re.sub(r"<[^>]+>", "", (d.get("summation") or [{}])[0].get("text", ""))
            salida.append({"reactome": e["stId"], "via": re.sub(r"<[^>]+>", "", e["name"]), "resumen": resumen[:1500]})
    return salida


_ABREVIATURAS = re.compile(r"\b(?:et al|e\.g|i\.e|vs|cf|Fig|Figs|Dr|Drs|Prof|No|Nr|ca|approx|bzw|ggf|vgl|Abb|Tab|"
                           r"z\.\s?B|d\.\s?h|u\.\s?a|p\.\s?ej|aprox|resp)\.$")


def _frases(texto, separadores=".;"):
    """Divide un texto en frases sin cortar en abreviaturas: «S. aureus», «et al.», «e.g.», «bzw.». Solo corta tras
    un punto (o punto y coma) seguido de espacio, si lo que sigue no empieza en minúscula y lo anterior no es una
    abreviatura conocida."""
    texto = re.sub(r"\s+", " ", texto or "").strip()
    frases, inicio = [], 0
    for m in re.finditer(rf"(?<=[{re.escape(separadores)}])\s+", texto):
        siguiente = texto[m.end():m.end() + 1]
        if siguiente.islower() or (texto[m.start() - 1] == "." and _ABREVIATURAS.search(texto[inicio:m.start()])):
            continue
        frases.append(texto[inicio:m.start()])
        inicio = m.end()
    if texto[inicio:]:
        frases.append(texto[inicio:])
    return frases


# Filtro de Latinoamérica para PubMed (países por MeSH y Costa Rica en el título o el resumen).
_GEO = '("Latin America"[MeSH] OR "Central America"[MeSH] OR "South America"[MeSH] OR "Mexico"[MeSH] OR ' \
       '"Caribbean Region"[MeSH] OR "Costa Rica"[tiab] OR "Latin America"[tiab])'


def _region(consulta, region):
    return f"({consulta}) AND {_GEO}" if region == "latam" else consulta


@_registrada
def europepmc(consulta, patron=None, maximo=60, citas_minimas=0, orden="relevancia"):
    """Busca en Europe PMC y devuelve frases de resúmenes que respaldan una afirmación.

    patron: expresión regular que debe cumplir la frase (p. ej., r"SREBP-?2.*LDL"). Por defecto ordena por
    relevancia; orden="citas" ordena por número de citas (favorece lo antiguo: no lo uses para buscar guías nuevas).
    Cada resultado trae DOI, PMCID, si es de acceso abierto y el tipo de publicación.
    """
    q = urllib.parse.quote(consulta)
    ordenar = "&sort=CITED%20desc" if orden == "citas" else ""
    r = json.loads(_get(f"https://www.ebi.ac.uk/europepmc/webservices/rest/search?query={q}&format=json"
                        f"&pageSize={maximo}&resultType=core{ordenar}"))
    salida = []
    for a in r["resultList"]["result"]:
        if (a.get("citedByCount") or 0) < citas_minimas:
            continue
        resumen = re.sub(r"<[^>]+>", "", a.get("abstractText", ""))
        frases = [f for f in _frases(resumen) if not patron or re.search(patron, f, re.I)]
        if patron and not frases:
            continue
        salida.append({"pmid": a.get("pmid"), "pmcid": a.get("pmcid"), "doi": a.get("doi"), "titulo": a["title"],
                       "autores": a.get("authorString"), "revista": a.get("journalTitle"), "anio": a.get("pubYear"),
                       "tipos": (a.get("pubTypeList") or {}).get("pubType", []),
                       "acceso_abierto": a.get("isOpenAccess") == "Y", "citas": a.get("citedByCount"),
                       "frases": frases[:2]})
    return salida


EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"
_AVISOS_PUBMED = {"RetractionIn": "retractado", "ErratumIn": "fe de erratas", "UpdateIn": "actualizado en",
                  "ExpressionOfConcernIn": "expresión de preocupación", "RepublishedIn": "republicado en",
                  "CorrectedandRepublishedIn": "corregido y republicado en"}


def _texto_xml(nodo):
    return re.sub(r"\s+", " ", "".join(nodo.itertext())).strip() if nodo is not None else ""


def _articulos_pubmed(xml):
    """Lee el XML de efetch (artículos y libros de PubMed): lo necesario para citar (autores, revista, año, DOI,
    PMCID), el tipo de publicación, el estado de indexación y los avisos de la NLM (retractado, fe de erratas,
    actualizado). Los identificadores de las referencias del artículo no se mezclan con los del propio artículo."""
    import xml.etree.ElementTree as ET
    salida = []
    for art in ET.fromstring(xml):
        if art.tag not in ("PubmedArticle", "PubmedBookArticle"):
            continue
        cita = art.find("MedlineCitation" if art.tag == "PubmedArticle" else "BookDocument")
        cita = art if cita is None else cita  # registros recortados (tests)
        fecha = cita.find(".//PubDate")
        anio = _texto_xml(fecha.find("Year")) if fecha is not None else ""
        if not anio and fecha is not None:
            anio = (re.match(r"\d{4}", _texto_xml(fecha.find("MedlineDate"))) or [""])[0]
        autores = []
        for autor in cita.findall(".//AuthorList/Author"):
            colectivo = _texto_xml(autor.find("CollectiveName"))
            autores.append(colectivo or f"{_texto_xml(autor.find('LastName'))} {_texto_xml(autor.find('Initials'))}"
                           .strip())
        datos = art.find("PubmedData" if art.tag == "PubmedArticle" else "PubmedBookData")
        lista = datos.find("ArticleIdList") if datos is not None else None
        ids = {}
        for identificador in (lista if lista is not None else []):
            ids.setdefault(identificador.get("IdType"), (identificador.text or "").strip())
        tipos = [_texto_xml(t) for t in cita.findall(".//PublicationTypeList/PublicationType")]
        avisos = {}
        for cc in cita.findall(".//CommentsCorrectionsList/CommentsCorrections"):
            if cc.get("RefType") in _AVISOS_PUBMED:
                referencia = _texto_xml(cc.find("PMID"))
                avisos.setdefault(_AVISOS_PUBMED[cc.get("RefType")], []).append(
                    f"PMID {referencia}" if referencia else _texto_xml(cc.find("RefSource")))
        avisos = [f"{tipo}: {', '.join(refs)}" for tipo, refs in avisos.items()]
        if "Retracted Publication" in tipos:
            avisos.insert(0, "RETRACTADO")
        if "Published Erratum" in tipos:
            avisos.append("es una fe de erratas")
        if "Retraction of Publication" in tipos:
            avisos.append("es un aviso de retracción")
        revisado = cita.find("ContributionDate")  # libros (LiverTox, LactMed): fecha de la última revisión
        resumen = " ".join(_texto_xml(t) for t in cita.findall(".//Abstract/AbstractText"))
        ensayos = [_texto_xml(n) for banco in cita.findall(".//DataBankList/DataBank")
                   if _texto_xml(banco.find("DataBankName")) == "ClinicalTrials.gov"
                   for n in banco.findall(".//AccessionNumber")]
        ensayos = list(dict.fromkeys(ensayos + re.findall(r"\bNCT\d{8}\b", resumen)))
        salida.append({
            "pmid": _texto_xml(cita.find("PMID")),
            "titulo": _texto_xml(cita.find(".//ArticleTitle")) or _texto_xml(cita.find(".//BookTitle")),
            "autores": autores,
            "revista": (_texto_xml(cita.find(".//ISOAbbreviation")) or _texto_xml(cita.find(".//Journal/Title"))
                        or _texto_xml(cita.find(".//BookTitle"))),
            "anio": anio, "doi": ids.get("doi"), "pmcid": ids.get("pmc"), "tipos": tipos,
            "estado": cita.get("Status") or ("libro" if art.tag == "PubmedBookArticle" else None),
            "avisos": avisos, "ensayos": ensayos,
            "revisado": "-".join(_texto_xml(revisado.find(c)).zfill(2) for c in ("Year", "Month", "Day"))
            if revisado is not None else None,
            "resumen": resumen,
        })
    return salida


@_registrada
def pubmed(consulta, patron=None, maximo=40, revisiones=False, completo=False, orden="relevancia", region=None):
    """Busca en PubMed (NCBI, NIH) y devuelve frases de resúmenes que respaldan una afirmación.

    Complementa a europepmc: PubMed ordena por relevancia e incluye revisiones y guías recientes.
    patron: expresión regular que debe cumplir la frase. revisiones=True limita a revisiones.
    orden="fecha" pone primero lo más reciente; region="latam" limita a Latinoamérica.
    Cada artículo trae lo necesario para citarlo (autores, revista, año, DOI, PMCID, tipo de publicación) y los
    avisos de la NLM: «RETRACTADO», «fe de erratas», «actualizado en». Revisa «avisos» antes de citar.
    """
    termino = _region(consulta, region) + (" AND review[pt]" if revisiones else "")
    q = urllib.parse.quote(termino)
    ids = json.loads(_get(f"{EUTILS}/esearch.fcgi?db=pubmed&term={q}&retmax={maximo}"
                          f"&sort={'pub_date' if orden == 'fecha' else 'relevance'}&retmode=json"))["esearchresult"]["idlist"]
    if not ids:
        return []
    xml = _get(f"{EUTILS}/efetch.fcgi?db=pubmed&id={','.join(ids)}&retmode=xml", timeout=120).decode("utf-8", "ignore")
    articulos = {a["pmid"]: a for a in _articulos_pubmed(xml)}
    salida = []
    for pmid in ids:  # en el orden de la búsqueda
        a = articulos.get(pmid)
        if not a:
            continue
        frases = [f for f in _frases(a.pop("resumen")) if not patron or re.search(patron, f, re.I)]
        if patron and not frases:
            continue
        a["frases"] = frases if completo else frases[:2]
        salida.append(a)
    return salida


@_registrada
def ncbi_gene(gen, organismo=9606):
    """Resumen curado del gen en NCBI Gene (NIH): función, nombres alternativos y localización."""
    q = urllib.parse.quote(f"{gen}[sym] AND {organismo}[taxid]")
    ids = json.loads(_get(f"{EUTILS}/esearch.fcgi?db=gene&term={q}&retmode=json"))["esearchresult"]["idlist"]
    if not ids:
        return None
    d = json.loads(_get(f"{EUTILS}/esummary.fcgi?db=gene&id={ids[0]}&retmode=json"))["result"][ids[0]]
    return {"gene_id": ids[0], "simbolo": d["name"], "nombre": d["description"], "alias": d.get("otheraliases"),
            "cromosoma": d.get("maplocation"), "resumen": d.get("summary")}


@_registrada
def nci_tesauro(nombre):
    """Definición del fármaco en el NCI Thesaurus (Instituto Nacional del Cáncer, NIH).

    Útil sobre todo en oncología: describe mecanismo, diana y clase con lenguaje revisado.
    """
    q = urllib.parse.quote(nombre)
    r = json.loads(_get(f"https://api-evsrest.nci.nih.gov/api/v1/concept/ncit/search?term={q}"
                        f"&type=match&include=definitions,synonyms&pageSize=3"))
    salida = []
    for c in r.get("concepts", []):
        defs = [d["definition"] for d in c.get("definitions", []) if d.get("source") in ("NCI", "NCI-GLOSS")]
        salida.append({"codigo": c["code"], "nombre": c["name"], "definiciones": defs})
    return salida


def _monografia_nih(nombre, editor, libro, patron=None):
    """Monografías del NIH indexadas en PubMed como libro (LiverTox, LactMed): texto completo del resumen y fecha
    de la última revisión («revisado»). Si el fármaco no tiene monografía propia, busca la de su clase."""
    import html as _html
    res = [r for r in pubmed(f'"{nombre}"[ti] AND "{editor}"[pb]', patron, maximo=5, completo=True)
           if libro in _html.unescape(r["revista"])]
    if not res:  # monografía de la clase («Sodium-Glucose Cotransporter-2 (SGLT2) Inhibitors»)
        res = [r for r in pubmed(f'"{nombre}" AND {libro.lower()}', patron, maximo=5, completo=True)
               if libro in _html.unescape(r["revista"])]
    for r in res:
        r["revista"] = _html.unescape(r["revista"])
    return res


@_registrada
def livertox(nombre, patron=None):
    """Hepatotoxicidad del fármaco según LiverTox (NIDDK, NIH), con PMID. nombre en inglés."""
    return _monografia_nih(nombre, "National Institute of Diabetes and Digestive and Kidney Diseases", "LiverTox", patron)


@_registrada
def lactmed(nombre, patron=None):
    """Paso a la leche y seguridad en la lactancia según LactMed (NICHD, NIH), con PMID. nombre en inglés."""
    return _monografia_nih(nombre, "National Institute of Child Health and Human Development", "LactMed", patron)


@_registrada
def medlineplus(nombre):
    """Información para pacientes en MedlinePlus (NLM, NIH), en español, a partir del RxCUI."""
    q = urllib.parse.quote(nombre)
    rx = json.loads(_get(f"https://rxnav.nlm.nih.gov/REST/rxcui.json?name={q}&search=2"))
    rxcui = (rx.get("idGroup", {}).get("rxnormId") or [None])[0]
    if not rxcui:
        return None
    r = json.loads(_get("https://connect.medlineplus.gov/service?mainSearchCriteria.v.cs=2.16.840.1.113883.6.88"
                        f"&mainSearchCriteria.v.c={rxcui}&informationRecipient.languageCode.c=es"
                        "&knowledgeResponseType=application/json"))
    entradas = r.get("feed", {}).get("entry", [])
    return [{"rxcui": rxcui, "titulo": e["title"]["_value"], "url": e["link"][0]["href"],
             "resumen": _texto_plano(e.get("summary", {}).get("_value", ""))[:800]} for e in entradas]


@_registrada
def chembl_actividad(nombre, diana=None, tipos=("IC50", "Ki", "Kd", "EC50"), maximo=200):
    """Potencia y selectividad del fármaco en ChEMBL: actividades medidas, agrupadas por diana.

    Devuelve por diana la mediana y el rango (nM) de cada tipo de medida y el número de ensayos.
    diana: filtra por texto en el nombre de la diana (p. ej., "HMG-CoA"). Útil para mostrar con
    números la afinidad por la diana frente a otras proteínas (selectividad).
    """
    from statistics import median
    info = chembl(nombre)
    if not info:
        return None
    base = "https://www.ebi.ac.uk/chembl/api/data"
    acts = json.loads(_get(f"{base}/activity.json?molecule_chembl_id={info[0]['chembl_id']}"
                           f"&standard_type__in={','.join(tipos)}&standard_units=nM&target_organism=Homo%20sapiens"
                           f"&limit={maximo}", timeout=120))["activities"]
    grupos = {}
    for a in acts:
        if a.get("standard_value") is None or (diana and diana.lower() not in (a.get("target_pref_name") or "").lower()):
            continue
        clave = (a["target_pref_name"], a["standard_type"])
        grupos.setdefault(clave, []).append((float(a["standard_value"]), a.get("document_chembl_id")))
    salida = [{"diana": d, "tipo": t, "mediana_nM": round(median(v for v, _ in vals), 3),
               "min_nM": min(v for v, _ in vals), "max_nM": max(v for v, _ in vals), "n": len(vals),
               "documentos": sorted({doc for _, doc in vals if doc})[:5]}
              for (d, t), vals in grupos.items()]
    return sorted(salida, key=lambda x: x["mediana_nM"])


@_registrada
def bindingdb(uniprot, smiles, corte_nM=100000):
    """Afinidades del fármaco por una diana en BindingDB (segunda fuente para las cifras de ChEMBL).

    Compara los ligandos de la diana con el SMILES del fármaco (de PubChem) sin estereoquímica.
    """
    from rdkit import Chem
    def clave(smi):
        m = Chem.MolFromSmiles(smi)
        if m is None:
            return None
        Chem.RemoveStereochemistry(m)
        return Chem.MolToSmiles(m)
    objetivo = clave(smiles)
    r = json.loads(_get(f"https://www.bindingdb.org/rest/getLigandsByUniprots?uniprot={uniprot}&cutoff={corte_nM}"
                        "&response=application/json", timeout=180))
    afinidades = r.get("getLindsByUniprotsResponse", r.get("getLigandsByUniprotsResponse", {})).get("affinities", [])
    return [{"tipo": a["affinity_type"], "valor_nM": a["affinity"], "pmid": a.get("pmid"), "doi": a.get("doi")}
            for a in afinidades if clave(a["smile"]) == objetivo]


@_registrada
def ema_epar(nombre, patron, documento="public-assessment-report", contexto=1):
    """Frases del informe público de evaluación de la EMA (EPAR) que cumplen un patrón, con su página.

    Solo existe para medicamentos autorizados por procedimiento centralizado (la mayoría de los nuevos,
    biológicos y oncológicos). nombre: marca comercial (p. ej., "enhertu"). El EPAR detalla
    farmacocinética, metabolitos, transportadores, exposición-respuesta y poblaciones especiales.
    """
    slug = re.sub(r"[^a-z0-9]+", "-", nombre.lower()).strip("-")
    destino = CACHE / f"epar-{slug}-{documento}.pdf"
    if not destino.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(_get(f"https://www.ema.europa.eu/en/documents/assessment-report/"
                                 f"{slug}-epar-{documento}_en.pdf", timeout=300))
    return _frases_pdf(destino, patron, contexto)


_EMA_JSON = "https://www.ema.europa.eu/en/documents/report/medicines-output-medicines_json-report_en.json"


@_registrada
def ema_medicamento(nombre):
    """Ficha de un medicamento de autorización centralizada en la EMA (datos públicos de la agencia, se renuevan
    cada día): estado, principio activo, ATC, indicación autorizada, fechas de autorización y de la última
    actualización, número de revisión y página del EPAR, que enlaza la información del producto (ficha técnica
    en inglés) y el informe de evaluación. nombre: marca comercial o principio activo en inglés."""
    ruta = CACHE / f"ema-medicamentos-{date.today().isoformat()}.json"
    if not ruta.exists():
        ruta.write_bytes(_get(_EMA_JSON, timeout=300))
    buscado = nombre.strip().lower()
    salida = []
    for m in json.loads(ruta.read_text(encoding="utf-8"))["data"]:
        if m.get("category") != "Human" or buscado not in (
                m.get("name_of_medicine", "").lower(), m.get("active_substance", "").lower(),
                m.get("international_non_proprietary_name_common_name", "").lower()):
            continue
        salida.append({"nombre": m["name_of_medicine"], "principio_activo": m.get("active_substance"),
                       "estado": m.get("medicine_status"), "procedimiento": m.get("ema_product_number"),
                       "atc": m.get("atc_code_human"), "titular": m.get(
                           "marketing_authorisation_developer_applicant_holder"),
                       "autorizacion": m.get("marketing_authorisation_date"),
                       "actualizado": m.get("last_updated_date"), "revision": m.get("revision_number"),
                       "generico": m.get("generic") == "Yes", "biosimilar": m.get("biosimilar") == "Yes",
                       "condicional": m.get("conditional_approval") == "Yes",
                       "indicacion": html.unescape(m.get("therapeutic_indication") or "")[:2000],
                       "url": m.get("medicine_url")})
    # el original antes que los genéricos y biosimilares
    return sorted(salida, key=lambda m: (m["generico"] or m["biosimilar"], m["estado"] != "Authorised"))


# Revistas donde se publican casi todos los ensayos pivotales (resultado principal de un ensayo de fase III).
_REVISTAS_ENSAYOS = {"N Engl J Med", "Lancet", "JAMA", "BMJ", "Ann Intern Med", "Lancet Oncol", "J Clin Oncol",
                     "Nat Med", "Lancet Diabetes Endocrinol", "Lancet Respir Med", "Lancet Infect Dis",
                     "Lancet Neurol", "Lancet HIV", "Eur Heart J", "Circulation", "Blood"}
_NO_PIVOTAL = re.compile(r"(?i)post[- ]hoc|secondary analys|subgroup|sub-?study|pooled|exploratory|"
                         r"prespecified analys|pre-specified analys|extension|cost-effectiveness|rationale and design|"
                         r"design and rationale|baseline characteristics|protocol")


@_registrada
def ensayos(farmaco, maximo=10, anios=None):
    """Ensayos clínicos aleatorizados con el fármaco en el título (PubMed), con su número de registro (NCT).
    Pone primero los de fase III con registro y deja detrás los análisis secundarios, los de subgrupos, las
    extensiones y los diseños. Para el ensayo pivotal de una indicación, compáralo con la sección 14 de la ficha
    de la FDA o la 5.1 de la ficha europea, que nombran los ensayos de la autorización; ensayo(NCT) da el registro.
    farmaco: nombre en inglés."""
    fecha = f' AND "last {anios} years"[dp]' if anios else ""
    consulta = (f'{farmaco}[ti] AND (randomized controlled trial[pt] OR "clinical trial, phase iii"[pt]){fecha} '
                f'NOT (review[pt] OR meta-analysis[pt] OR letter[pt] OR comment[pt] OR editorial[pt])')
    resultados = pubmed(consulta, maximo=max(40, 3 * maximo))
    for orden, r in enumerate(resultados):
        r["puntos"] = (2 * ("Clinical Trial, Phase III" in r["tipos"]) + bool(r["ensayos"])
                       + ("Multicenter Study" in r["tipos"]) + 2 * (r["revista"] in _REVISTAS_ENSAYOS)
                       - 3 * bool(_NO_PIVOTAL.search(r["titulo"]))
                       - 10 * any(a.startswith("RETRACTADO") for a in r["avisos"]))
        r["_orden"] = orden
    resultados.sort(key=lambda r: (-r["puntos"], r["_orden"]))  # a igualdad, el orden de relevancia de PubMed
    return [{k: v for k, v in r.items() if k != "_orden"} for r in resultados[:maximo]]


@_registrada
def ensayo(nct):
    """Registro de un ensayo en ClinicalTrials.gov: título, fase, estado, tamaño, variable principal, fechas y si
    tiene resultados publicados en el registro."""
    d = json.loads(_get(f"https://clinicaltrials.gov/api/v2/studies/{nct}"))
    p = d.get("protocolSection", {})
    ident, estado, diseno = p.get("identificationModule", {}), p.get("statusModule", {}), p.get("designModule", {})
    return {"nct": nct, "titulo": ident.get("officialTitle") or ident.get("briefTitle"),
            "acronimo": ident.get("acronym"), "fases": diseno.get("phases"),
            "estado": estado.get("overallStatus"), "inicio": (estado.get("startDateStruct") or {}).get("date"),
            "fin": (estado.get("primaryCompletionDateStruct") or {}).get("date"),
            "participantes": (diseno.get("enrollmentInfo") or {}).get("count"),
            "variables_principales": [o.get("measure") for o in
                                      p.get("outcomesModule", {}).get("primaryOutcomes", [])],
            "patrocinador": p.get("sponsorCollaboratorsModule", {}).get("leadSponsor", {}).get("name"),
            "con_resultados": bool(d.get("hasResults")), "url": f"https://clinicaltrials.gov/study/{nct}"}


def _frases_pdf(ruta, patron, contexto=1, maximo=None):
    """Frases de un PDF que cumplen el patrón (regex, sin distinguir mayúsculas), con su página y las
    `contexto - 1` frases vecinas a cada lado."""
    import pymupdf
    salida = []
    for n, pagina_ in enumerate(pymupdf.open(ruta), 1):
        frases = _frases(pagina_.get_text())
        for i, f in enumerate(frases):
            if re.search(patron, f, re.I):
                salida.append({"pagina": n, "texto": " ".join(frases[max(0, i - contexto + 1):i + contexto])})
                if maximo and len(salida) >= maximo:
                    return salida
    return salida


@_registrada
def pdf_texto(url, patrones=(), contexto=1, maximo=8):
    """Frases de un PDF público (guía nacional, documento de posición, informe técnico) que cumplen cada patrón,
    con el número de página para citarlas. url: dirección del PDF o ruta local. El PDF se guarda en la caché.
    Si el PDF es una imagen escaneada no tiene texto: lo indica en «aviso». Si la dirección no devuelve un PDF
    (página de acceso, aviso de cookies), lanza ValueError: busca el enlace con pdf_desde_pagina."""
    import hashlib
    import pymupdf
    if Path(url).exists():
        ruta = Path(url)
    else:
        ruta = CACHE / f"pdf-{hashlib.sha1(url.encode()).hexdigest()[:16]}.pdf"
        if not ruta.exists():
            contenido = _get(url, timeout=300, cabeceras=_NAVEGADOR)
            if b"%PDF" not in contenido[:1024]:
                raise ValueError(f"{url} no devolvió un PDF (¿página del repositorio o de la editorial?)")
            ruta.write_bytes(contenido)
    documento = pymupdf.open(ruta)
    caracteres = sum(len(p.get_text()) for p in documento)
    return {"url": url, "titulo": (documento.metadata or {}).get("title") or None, "paginas": len(documento),
            "aviso": "sin texto extraíble: puede ser un escaneo" if caracteres < 20 * len(documento) else None,
            "fragmentos": {p: _frases_pdf(ruta, p, contexto, maximo) for p in patrones}}


def _enlace_pdf(contenido, base):
    """Enlace al PDF en la página de un artículo o de un repositorio: metadato citation_pdf_url (editoriales, OJS,
    EPrints), descargas de DSpace (bitstreams) o el primer enlace .pdf."""
    candidatos = re.findall(r'<meta[^>]+name="citation_pdf_url"[^>]+content="([^"]+)"', contenido)
    candidatos += re.findall(r'href="([^"]*/bitstreams?/[^"]+?/(?:download|content)[^"]*)"', contenido)
    candidatos += re.findall(r'href="([^"]*/bitstream/[^"]+?\.pdf[^"]*)"', contenido)
    candidatos += re.findall(r'href="([^"]+?\.pdf(?:\?[^"]*)?)"', contenido)
    for candidato in candidatos:
        url = urllib.parse.urljoin(base, html.unescape(candidato))
        anfitrion = urllib.parse.urlsplit(url).netloc
        if url.startswith("http") and not (anfitrion.endswith(".local") or ".local:" in anfitrion
                                           or anfitrion.startswith("localhost")):
            return url
    return None


@_registrada
def pdf_desde_pagina(url):
    """Dirección del PDF enlazado desde la página de un artículo o de un repositorio institucional, o None."""
    contenido, final = _get(url, cabeceras=_NAVEGADOR, con_url=True)
    return _enlace_pdf(contenido.decode("utf-8", "ignore"), final)


UNPAYWALL_EMAIL = os.environ.get("UNPAYWALL_EMAIL", "skills-docencia@example.org")


@_registrada
def acceso_abierto(doi, patrones=(), contexto=1, maximo=8):
    """Copias legales en acceso abierto de un artículo (Unpaywall: repositorios institucionales, la editorial),
    para cuando PMC o la editorial no dejan descargar el texto completo. Lee la primera que se pueda descargar
    como PDF y devuelve dónde se leyó, la versión (publicada, aceptada…), la licencia y los fragmentos con su
    página. Cita siempre la versión publicada (revista, DOI); la copia solo sirve para leer y verificar."""
    try:
        d = json.loads(_get(f"https://api.unpaywall.org/v2/{urllib.parse.quote(doi)}"
                            f"?email={urllib.parse.quote(UNPAYWALL_EMAIL)}"))
    except RuntimeError as e:
        if "404" not in str(e):
            raise
        return {"doi": doi, "texto_completo": False, "es_acceso_abierto": None, "intentos": [],
                "detalle": "Unpaywall aún no tiene este DOI (artículo muy reciente o DOI incorrecto)"}
    intentos = []
    for ubicacion in d.get("oa_locations") or []:
        for candidato in (ubicacion.get("url_for_pdf"), ubicacion.get("url_for_landing_page"), ubicacion.get("url")):
            if not candidato or candidato in [i["url"] for i in intentos]:
                continue
            if re.search(r"ncbi\.nlm\.nih\.gov/pmc|europepmc\.org", candidato):
                continue  # PMC se intenta con pmc_texto
            try:
                try:
                    url_pdf, r = candidato, pdf_texto(candidato, patrones, contexto, maximo)
                except ValueError:  # no es un PDF: buscar el enlace en la página
                    url_pdf = pdf_desde_pagina(candidato)
                    if not url_pdf:
                        raise ValueError("la página no enlaza un PDF")
                    r = pdf_texto(url_pdf, patrones, contexto, maximo)
            except Exception as e:  # noqa: BLE001
                intentos.append({"url": candidato, "error": str(e)[:200]})
                continue
            r.update({"doi": doi, "texto_completo": True, "fuente": url_pdf, "pagina_origen": candidato,
                      "tipo": ubicacion.get("host_type"), "version": ubicacion.get("version"),
                      "licencia": ubicacion.get("license")})
            return r
    return {"doi": doi, "texto_completo": False, "es_acceso_abierto": d.get("is_oa"), "intentos": intentos}


@_registrada
def ids(*identificadores):
    """PMID, PMCID y DOI de uno o varios artículos (conversor de identificadores de PMC); acepta cualquiera de
    los tres. Si el artículo no está en PMC, el conversor no da su DOI: búscalo con pubmed("<PMID>[uid]")."""
    lista = ",".join(urllib.parse.quote(str(i).strip()) for i in identificadores)
    d = json.loads(_get(f"https://pmc.ncbi.nlm.nih.gov/tools/idconv/api/v1/articles/?format=json&ids={lista}"))
    return [{"pmid": str(r["pmid"]) if r.get("pmid") else None, "pmcid": r.get("pmcid"), "doi": r.get("doi"),
             "pedido": r.get("requested-id")} for r in d.get("records", [])]


def pmcid(pmid):
    """PMCID de un artículo de PubMed, o None si no está en PMC."""
    registros = ids(pmid)
    return registros[0]["pmcid"] if registros else None


def _bloques_jats(xml):
    """Bloques de texto de un artículo JATS (PMC, Europe PMC) con la sección o la tabla de la que salen, para citar
    «tabla 2» o «Results» además de la frase."""
    import xml.etree.ElementTree as ET
    raiz = ET.fromstring(re.sub(r"<!DOCTYPE[^>]*>", "", xml))
    articulo = raiz if raiz.tag == "article" else raiz.find(".//article")
    if articulo is None:
        raise ValueError("sin elemento article")

    def nombre(nodo):
        return nodo.tag.split("}")[-1] if isinstance(nodo.tag, str) else ""

    bloques = [("Resumen", _texto_xml(r)) for r in articulo.iter("abstract")]

    def recorrer(nodo, ruta):
        for hijo in nodo:
            n = nombre(hijo)
            if n == "sec":
                titulo = _texto_xml(hijo.find("title"))
                recorrer(hijo, ruta + ([titulo] if titulo else []))
            elif n in ("table-wrap", "fig", "boxed-text"):
                etiqueta = " ".join(x for x in (_texto_xml(hijo.find("label")),
                                                _texto_xml(hijo.find("caption/title"))) if x)
                bloques.append((etiqueta[:90] or " › ".join(ruta) or n, _texto_xml(hijo)))
            elif n in ("p", "list", "disp-quote", "def-list", "statement"):
                bloques.append((" › ".join(ruta) or "Texto", _texto_xml(hijo)))
            elif n != "title":
                recorrer(hijo, ruta)

    for parte in ("body", "floats-group"):
        if articulo.find(parte) is not None:
            recorrer(articulo.find(parte), [])
    return [(etiqueta, texto) for etiqueta, texto in bloques if texto]


@_registrada
def pmc_texto(pmcid, patrones=(), contexto=350, maximo=6):
    """Texto completo de un artículo de PMC cuando la editorial lo permite (p. ej., acceso abierto).

    Devuelve, para cada patrón (regex, sin distinguir mayúsculas), hasta `maximo` fragmentos con la sección o la
    tabla de la que salen. Lo pide a NCBI y, si NCBI no lo da (la editorial no permite descargarlo, o el servicio
    falla), a Europe PMC; «fuente» dice cuál respondió. Si ninguno lo tiene, lo indica: prueba entonces
    texto_completo (busca copias en acceso abierto) o usa el resumen de PubMed.
    """
    numero = str(pmcid).upper().replace("PMC", "")
    intentos = (("NCBI PMC", f"{EUTILS}/efetch.fcgi?db=pmc&id={numero}&retmode=xml"),
                ("Europe PMC", f"https://www.ebi.ac.uk/europepmc/webservices/rest/PMC{numero}/fullTextXML"))
    errores = []
    for fuente, url in intentos:
        try:
            xml = _get(url, timeout=180, intentos=2).decode("utf-8", "ignore")
        except RuntimeError as e:
            errores.append(str(e))
            continue
        if "does not allow downloading of the full text" in xml or "<body" not in xml:
            errores.append(f"{fuente}: sin texto completo")
            continue
        try:
            bloques = _bloques_jats(xml)
        except Exception:  # noqa: BLE001  (XML con entidades no estándar: texto plano, sin secciones)
            bloques = [(None, re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", xml))))]
        fragmentos = {}
        for p in patrones:
            fragmentos[p] = [{"seccion": etiqueta, "texto": texto[max(0, m.start() - contexto):m.end() + contexto]}
                             for etiqueta, texto in bloques for m in re.finditer(p, texto, re.I)][:maximo]
        return {"pmcid": f"PMC{numero}", "texto_completo": True, "fuente": fuente,
                "caracteres": sum(len(t) for _, t in bloques), "fragmentos": fragmentos}
    return {"pmcid": f"PMC{numero}", "texto_completo": False, "detalle": "; ".join(errores)}


@_registrada
def texto_completo(identificador, patrones=(), maximo=6):
    """Texto completo de un artículo por la vía que funcione: PMC (NCBI y Europe PMC) y, si no, copias legales en
    acceso abierto (Unpaywall). identificador: PMID, PMCID o DOI. Los fragmentos llevan su sección (PMC) o su
    página (PDF); «via» dice de dónde salió el texto."""
    identificador = str(identificador).strip()
    info = {"doi": identificador} if re.match(r"10\.\d{4,}/", identificador) else (
        {"pmcid": identificador.upper()} if identificador.upper().startswith("PMC") else {"pmid": identificador})
    try:
        if "pmid" in info:
            articulo = pubmed(f"{info['pmid']}[uid]", maximo=1)
            if articulo:
                info.update({k: articulo[0][k] for k in ("pmcid", "doi") if articulo[0].get(k)})
        else:
            for r in ids(info.get("pmcid") or info.get("doi")):
                info.update({k: v for k, v in r.items() if v and k in ("pmid", "pmcid", "doi")})
    except RuntimeError:
        pass
    detalles = []
    if info.get("pmcid"):
        r = pmc_texto(info["pmcid"], patrones, maximo=maximo)
        if r["texto_completo"]:
            return {**info, **r, "via": r["fuente"]}
        detalles.append(r.get("detalle"))
    if info.get("doi"):
        r = acceso_abierto(info["doi"], patrones, maximo=maximo)
        if r["texto_completo"]:
            return {**info, **r, "via": f"acceso abierto ({r.get('tipo')}, {r.get('version')})"}
        detalles.append(r.get("detalle") or f"Unpaywall: {len(r.get('intentos', []))} copias, ninguna descargable")
    return {**info, "texto_completo": False, "detalle": "; ".join(d for d in detalles if d) or "sin PMCID ni DOI"}


@_registrada
def openfda(nombre, secciones=("clinical_pharmacology", "mechanism_of_action", "microbiology", "pharmacokinetics",
                               "drug_interactions", "pharmacogenomics", "use_in_specific_populations"), largo=2500):
    """Ficha de la FDA ya separada por secciones (openFDA). nombre: genérico en inglés.
    En los antimicrobianos, 12.1 suele remitir a 12.4 («microbiology»): ahí están el mecanismo y la resistencia.
    Las fichas antiguas sin formato PLR no tienen ese campo; su «Microbiology» va dentro de clinical_pharmacology."""
    q = urllib.parse.quote(f'openfda.generic_name:"{nombre}"')
    fichas = json.loads(_get(f"https://api.fda.gov/drug/label.json?search={q}&limit=25", timeout=90))["results"]
    r = max(fichas, key=lambda f: _puntuar_ficha(f, nombre, secciones))
    salida = {"set_id": r.get("set_id"), "fecha": r.get("effective_time"),
              "producto": "; ".join(r.get("openfda", {}).get("brand_name", []) + r.get("openfda", {}).get(
                  "generic_name", [])) or None}
    for s in secciones:
        if r.get(s):
            salida[s] = " ".join(r[s])[:largo]
    return salida


def _puntuar_ficha(ficha, nombre, secciones):
    """Orden de preferencia entre fichas de openFDA: primero la del principio activo solo (no una combinación),
    después la que tiene más secciones de las pedidas y, a igualdad, la más reciente."""
    genericos = [g.strip().upper() for g in ficha.get("openfda", {}).get("generic_name", [])]
    buscado = nombre.strip().upper()
    solo = any(g.startswith(buscado) and not re.search(r" AND |,|/|;", g) for g in genericos)  # admite la sal
    return solo, sum(bool(ficha.get(s)) for s in secciones), ficha.get("effective_time", "")


@_registrada
def openfda_eventos(nombre, maximo=15):
    """Reacciones más notificadas en FAERS (FDA). Son notificaciones espontáneas: no dan frecuencias
    ni causalidad; sirven solo para orientar qué buscar en la ficha técnica."""
    q = urllib.parse.quote(f'patient.drug.openfda.generic_name:"{nombre}"')
    r = json.loads(_get(f"https://api.fda.gov/drug/event.json?search={q}"
                        f"&count=patient.reaction.reactionmeddrapt.exact&limit={maximo}", timeout=90))
    return r["results"]


@_registrada
def cpic(nombre):
    """Pares gen-fármaco de CPIC (farmacogenética) con nivel de evidencia y recomendaciones.

    Niveles A y B: hay recomendación de prescripción. Devuelve los pares A/B con sus recomendaciones
    por fenotipo, y la lista de pares C/D sin recomendación. nombre: genérico en inglés, en minúsculas.
    """
    base = "https://api.cpicpgx.org/v1"
    q = urllib.parse.quote(nombre.lower())
    pares = json.loads(_get(f"{base}/pair_view?drugname=eq.{q}&select=genesymbol,cpiclevel,guidelinename,"
                            "guidelineurl,clinpgxlevel,pmids,usedforrecommendation"))
    recs = json.loads(_get(f"{base}/recommendation_view?drugname=eq.{q}&select=lookupkey,drugrecommendation,"
                           "classification,implications,population"))
    fuertes = [p for p in pares if p.get("cpiclevel") in ("A", "B")]
    for p in fuertes:
        p["recomendaciones"] = [r for r in recs if p["genesymbol"] in (r.get("lookupkey") or {})]
    return {"con_recomendacion": fuertes,
            "sin_recomendacion": [(p["genesymbol"], p["cpiclevel"]) for p in pares if p not in fuertes]}


@_registrada
def fda_indicaciones(nombre, ultimas=6, cartas=True):
    """Indicaciones aprobadas por la FDA: texto vigente de la ficha y aprobaciones recientes con fecha.

    - Ficha vigente (openFDA): sección 1 completa, fecha y «Recent Major Changes».
    - Historial (Drugs@FDA): suplementos de eficacia aprobados (nuevas indicaciones o cambios de uso),
      del más reciente al más antiguo, con la frase de la carta de aprobación que dice qué se aprobó.
    Usa la solicitud del innovador (NDA o BLA), no las de genéricos (ANDA). nombre: genérico en inglés.
    """
    import pymupdf
    palabras = [w for w in re.split(r"[\s-]+", nombre.lower()) if len(w) > 2]
    q = "+AND+".join(f"openfda.generic_name:{urllib.parse.quote(w)}" for w in palabras)
    apps = json.loads(_get(f"https://api.fda.gov/drug/drugsfda.json?search={q}&limit=50", timeout=90))["results"]
    apps = [a for a in apps if a["application_number"].startswith(("NDA", "BLA"))]
    if not apps:
        return None
    # El innovador suele ser la solicitud con más suplementos de eficacia.
    app = max(apps, key=lambda a: sum(s.get("submission_class_code") == "EFFICACY" for s in a.get("submissions", [])))
    subs = sorted((s for s in app.get("submissions", []) if s.get("submission_status") == "AP"
                   and (s.get("submission_class_code") == "EFFICACY" or s.get("submission_type") == "ORIG")),
                  key=lambda s: s["submission_status_date"], reverse=True)
    historial, vistas = [], set()
    for sub in subs[:ultimas]:
        docs = {d["type"]: d["url"] for d in sub.get("application_docs", [])}
        fecha = sub["submission_status_date"]
        item = {"fecha": f"{fecha[:4]}-{fecha[4:6]}-{fecha[6:]}", "tipo": sub["submission_type"],
                "numero": sub["submission_number"], "carta": docs.get("Letter"), "ficha": docs.get("Label")}
        if cartas and item["carta"] and (item["carta"], item["numero"]) not in vistas:
            vistas.add((item["carta"], item["numero"]))
            try:
                destino = CACHE / ("fda-" + re.sub(r"[^\w.]+", "_", item["carta"].rsplit("/", 1)[-1]))
                if not destino.exists():
                    CACHE.mkdir(parents=True, exist_ok=True)
                    destino.write_bytes(_get(item["carta"].replace(" ", "%20"), timeout=120,
                                             cabeceras={"User-Agent": "Mozilla/5.0"}))
                texto = re.sub(r"\s+", " ", " ".join(p.get_text() for p in pymupdf.open(destino)))
                fin = r"(?=Prior Approval|We have completed|APPROVAL &|This information will|CONTENT OF LABELING|$)"
                verbo = r"(?:provides?|provided|proposes)(?: for)? (.{20,900}?)"
                m = (re.search(rf"S-?0*{int(item['numero'])}\b[^.]{{0,40}}?{verbo}{fin}", texto)
                     or re.search(verbo + fin, texto)
                     or re.search(r"is indicated (for .{20,600}?\.)", texto))
                item["aprobado"] = m.group(1).strip() if m else None
            except Exception as e:  # noqa: BLE001
                item["aprobado"] = f"(no se pudo leer la carta: {e})"
        historial.append(item)
    ficha = {}
    try:
        etiqueta = json.loads(_get(f"https://api.fda.gov/drug/label.json?search={q}+AND+openfda.application_number:"
                                   f"{app['application_number']}&sort=effective_time:desc&limit=1", timeout=90))
        r = etiqueta["results"][0]
        ficha = {"fecha": r.get("effective_time"), "cambios_recientes": " ".join(r.get("recent_major_changes", [])),
                 "indicaciones": " ".join(r.get("indications_and_usage", []))}
    except Exception as e:  # noqa: BLE001
        ficha = {"error": str(e)}
    return {"solicitud": app["application_number"], "titular": app.get("sponsor_name"),
            "marcas": sorted({p.get("brand_name") for p in app.get("products", [])}),
            "ficha_vigente": ficha, "aprobaciones_recientes": historial}


# --- Enfermedades: definiciones, ontologías y guías ------------------------------------

@_registrada
def mesh(termino):
    """Definición curada de MeSH (NLM) y su identificador. termino: en inglés.
    «mesh» es el identificador que se cita (D…, p. ej., D003924); «uid» es el número interno de Entrez.
    «sinonimos» son los términos de entrada de MeSH: úsalos con OR en guias y pubmed para no perder registros."""
    base = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"
    q = urllib.parse.quote(termino)
    ids = json.loads(_get(f"{base}/esearch.fcgi?db=mesh&term={q}&retmode=json"))["esearchresult"]["idlist"][:3]
    salida = []
    for uid in ids:
        r = json.loads(_get(f"{base}/esummary.fcgi?db=mesh&id={uid}&retmode=json"))["result"][uid]
        salida.append({"mesh": r.get("ds_meshui"), "uid": uid, "termino": (r.get("ds_meshterms") or [""])[0],
                       "sinonimos": (r.get("ds_meshterms") or [])[1:], "definicion": r.get("ds_scopenote", "").strip()})
    return salida


@_registrada
def mondo(termino):
    """Enfermedad en la ontología MONDO (EBI OLS): identificador, definición y sinónimos. termino: en inglés."""
    q = urllib.parse.quote(termino)
    d = json.loads(_get(f"https://www.ebi.ac.uk/ols4/api/search?q={q}&ontology=mondo&rows=3"))
    return [{"id": x.get("obo_id"), "nombre": x.get("label"), "definicion": " ".join(x.get("description") or [])}
            for x in d["response"]["docs"]]


_NO_GUIA = ('letter[pt] OR comment[pt] OR editorial[pt] OR "published erratum"[pt] OR "retracted publication"[pt] '
            'OR "retraction of publication"[pt]')
_PALABRAS_GUIA = ('guideline[ti] OR guidelines[ti] OR "consensus statement"[ti] OR "consensus report"[ti] OR '
                  '"position statement"[ti] OR "position paper"[ti] OR "standards of care"[ti] OR recommendations[ti]')
_TITULO_GUIA = re.compile(r"(?i)\bguidelines?\b|consensus (statement|report)|position (statement|paper)|"
                          r"standards of care|recommendations|leitlinie|gu[ií]a")
_COMENTARIO = re.compile(r"(?i)\b(adherence|compliance|survey|questionnaire|dissenting|revisiting|what'?s next|"
                         r"highlights|implementation|comparing|comparison|comparative|appraisal|quality assessment|"
                         r"concordance|alignment|knowledge|awareness|perspectives?|attitudes?|misinformation|"
                         r"controvers\w*|commentary|correction|erratum|did not endorse|turning point|impact of|"
                         r"performance of|chatgpt|deepseek|large language|for whom|limits of|lessons|insights|"
                         r"reflections?|critique|outcomes of|under new|scoping review|need for)\b|^\[|\?|"
                         r"^(?:a|an) .*\b(?:study|trial|assay|analysis|cohort)\b")
# «Guía» como núcleo del título (no una mención): hasta 8 palabras desde el inicio o tras «:», y luego la palabra de
# guía seguida de for/on/of/in/to/by, de «:» o del final. «Korean Guidelines for…», «SSC: international guidelines
# for…», «…: a position paper of…»; no «… and sepsis guidelines and the limits…».
_NUCLEO_GUIA = re.compile(r"(?i)(?:^|[:.]\s+)(?:[\w/.,'’()-]+\s+){0,8}?(?:guidelines?|recommendations|"
                          r"consensus (?:statement|report)|position (?:statement|paper)|standards of care)"
                          r"(?:\s+(?:for|on|of|in|to|by)\b|\s*:|\s*$)")
_TIPOS_GUIA = {"Practice Guideline", "Guideline", "Consensus Statement", "Consensus Development Conference",
               "Consensus Development Conference, NIH"}
# Títulos de estudios y revisiones que mencionan guías sin serlo («…: A Narrative Review of Evidence and
# Recommendations», «…: a case report», «High Prevalence of …: Key Predictors and Screening Recommendations»).
_DISENO = re.compile(r"(?i)\b(?:narrative|systematic|literature|integrative|umbrella|rapid) review\b|\breview of\b|"
                     r"\bcase (?:report|series)\b|\ba case\b|\bprevalence\b|\bpredictors?\b|\beffects? of\b|"
                     r"\bsimulation\b|\bvalidation of\b|health technology assessment|development process|research synthesis|"
                     r"recent advances|\ban update on\b|\bto (?:\w+ )?guidelines\b|\bmeta-analysis\b|"
                     r"\b(?:delphi|cohort|cross-sectional|retrospective|prospective|qualitative|observational|"
                     r"randomi[sz]ed)(?: \w+)? (?:study|analysis|trial)\b")
# Alcance general (manejo, diagnóstico, tratamiento) frente a un aspecto concreto (dieta, tabaco, un dispositivo).
_ALCANCE = re.compile(r"(?i)\b(?:management|managing|treatment|diagnosis|standards of care|algorithm)\b")
_SERIE_NORMAS = '"standards of care"[ti] OR "standards of medical care"[ti]'
_NO_DISTINTIVAS = set("type disease diseases syndrome disorder disorders infection infections chronic acute severe "
                      "adult adults children patients mellitus".split())


def _en_titulo(expresion):
    """«a b OR c d» → «(a[ti] AND b[ti]) OR (c[ti] AND d[ti])»: cada alternativa, con todas sus palabras en el título."""
    alternativas = [a for a in re.split(r"\s+OR\s+", expresion.strip().strip("()")) if a.strip()]
    return " OR ".join("(" + " AND ".join(f"{p}[ti]" for p in re.findall(r"[\w-]+", a)) + ")" for a in alternativas)


def _puntuar_guia(r, enfermedad):
    """Cuánto parece una guía general sobre la enfermedad: el tipo «Practice Guideline» o la guía como núcleo del
    título (las aún sin indexar no tienen tipo), la enfermedad en el título, el alcance general (manejo,
    diagnóstico, tratamiento) y las normas por capítulos; resta si parece un comentario, un estudio o una
    revisión sobre guías, o si está retractada."""
    titulo = r["titulo"].rstrip(". ")
    tipo_guia = bool(_TIPOS_GUIA & set(r.get("tipos", [])))
    puntos = 3 if (tipo_guia or _NUCLEO_GUIA.search(titulo)) else 0
    puntos += 1 if _TITULO_GUIA.search(r["titulo"]) else 0
    alternativas = [re.findall(r"[\w-]+", a.lower()) for a in re.split(r"\s+OR\s+", enfermedad)]
    puntos += 2 if any(a and all(p in r["titulo"].lower() for p in a) for a in alternativas) else 0
    puntos += 1 if _ALCANCE.search(titulo) else 0
    puntos += 2 if len(r.get("secciones", [])) >= 2 else 0  # normas publicadas por capítulos: guía completa
    puntos -= 4 if _COMENTARIO.search(r["titulo"]) else 0
    puntos -= 3 if not tipo_guia and _DISENO.search(titulo) else 0
    puntos -= 10 if any(a.startswith("RETRACTADO") for a in r.get("avisos", [])) else 0
    return puntos


def _serie(titulo):
    """Nombre de la serie de una guía publicada por capítulos, para agruparlos: la parte del título (antes o
    después de «:») con la palabra de guía y el año («Standards of Care in Diabetes-2026», «ISPAD Clinical
    Practice Consensus Guidelines 2022»). None si el título no tiene «:»."""
    partes = titulo.rstrip(". ").split(":")
    if len(partes) < 2:
        return None
    for parte in partes:
        if re.search(r"(?:19|20)\d{2}\b", parte) and _TITULO_GUIA.search(parte):
            return parte.strip()
    return None


def _distintivas(enfermedad):
    """Palabras propias de la enfermedad («diabetes» en «type 2 diabetes»): las normas generales de una sociedad
    la nombran sin el subtipo («Standards of Care in Diabetes-2026»)."""
    return sorted({p for p in re.findall(r"[^\W\d_][\w-]+", enfermedad.lower())
                   if len(p) >= 5 and p not in _NO_DISTINTIVAS})


@_registrada
def guias(enfermedad, maximo=10, titulo=False, anios=5, region=None):
    """Guías de práctica clínica y consensos recientes en PubMed, de la más pertinente y reciente a la menos.

    Busca en tres pasadas: las indexadas como guía («Practice Guideline», que PubMed asigna meses después de
    publicarse y nunca en revistas fuera de MEDLINE); las aún sin indexar con palabras de guía en el título, para
    que aparezcan las guías de los últimos meses; y las normas anuales por capítulos («Standards of Care in
    Diabetes-2026»), que no llevan el tipo de guía y nombran la enfermedad sin el subtipo. Excluye cartas,
    editoriales, fe de erratas y retractaciones, agrupa las copublicaciones (misma guía en varias revistas) y los
    capítulos de una misma serie («secciones»; de una serie anual queda la más reciente y las demás van en
    «anteriores»), y marca los avisos de la NLM. Cada resultado lleva PMCID y DOI: léelo con pmc_texto o
    texto_completo.
    enfermedad: en inglés; varias formas con OR («hospital-acquired pneumonia OR ventilator-associated pneumonia»).
    titulo=True exige la enfermedad en el título también en la primera pasada (menos ruido).
    """
    fecha = f'"last {anios} years"[dp]'
    tema = f"({_en_titulo(enfermedad)})" if titulo else f"({enfermedad})"
    pasadas = [
        ("indexada", f'{tema} AND (practice guideline[pt] OR guideline[pt] OR consensus[ti] OR '
                     f'"standards of care"[ti]) AND {fecha} NOT ({_NO_GUIA})'),
        ("sin indexar", f'({_en_titulo(enfermedad)}) AND ({_PALABRAS_GUIA}) AND (inprocess[sb] OR publisher[sb] OR '
                        f'pubmednotmedline[sb]) AND {fecha} NOT ({_NO_GUIA})'),
    ]
    if _distintivas(enfermedad):
        nucleo = " OR ".join(f"{p}[ti]" for p in _distintivas(enfermedad))
        pasadas.append(("normas", f'({nucleo}) AND ({_SERIE_NORMAS}) AND "last 2 years"[dp] NOT ({_NO_GUIA})'))
    resultados, por_titulo, por_serie = {}, {}, {}
    for pasada, consulta in pasadas:
        for r in pubmed(consulta, maximo=max(100, 5 * maximo) if pasada == "indexada" else max(20, 3 * maximo),
                        orden="relevancia" if pasada == "indexada" else
                        "fecha", region=region):
            if r["pmid"] in resultados:
                continue
            r["pasada"] = pasada
            clave = re.sub(r"\W+", "", r["titulo"].lower())
            if clave in por_titulo:  # la misma guía publicada en varias revistas
                por_titulo[clave].setdefault("copublicaciones", []).append(f"{r['pmid']} ({r['revista']})")
                continue
            serie = _serie(r["titulo"])
            clave_serie = re.sub(r"\W+", "", (serie or "").lower())
            if serie and clave_serie in por_serie:  # otro capítulo de las mismas normas
                primero = por_serie[clave_serie]
                primero["serie"] = serie
                primero.setdefault("secciones", []).append(f"{r['pmid']}: {r['titulo']}")
                continue
            por_titulo[clave] = resultados[r["pmid"]] = r
            if serie:
                por_serie[clave_serie] = r
    ultimas = {}  # de una serie anual (normas de 2026 y de 2025), solo la más reciente
    for r in sorted(list(resultados.values()), key=lambda r: r["anio"] or "", reverse=True):
        base = re.sub(r"\W+|(?:19|20)\d{2}", "", (_serie(r["titulo"]) or "").lower())
        if len(base) < 15:
            continue
        if base in ultimas:
            ultimas[base].setdefault("anteriores", []).append(f"{r['pmid']}: {_serie(r['titulo'])}")
            del resultados[r["pmid"]]
        else:
            ultimas[base] = r
    for r in resultados.values():
        r["puntos"] = _puntuar_guia(r, enfermedad)
    return sorted(resultados.values(), key=lambda r: (r["puntos"], r["anio"] or ""), reverse=True)[:maximo]


_GENERICAS = set("guideline guidelines international clinical practice management treatment diagnosis therapy "
                 "recommendations recommendation update updated consensus statement report position paper adults "
                 "adult patients patient care standards evidence based summary executive society association "
                 "american european national guidance".split())
_VACIAS = set("the of and for in on with to a an by from at as or its their de del la el los las y en para con "
              "por".split())


def _palabras_titulo(titulo):
    return {p for p in re.findall(r"[^\W\d_]+", (titulo or "").lower()) if len(p) > 2 and p not in _VACIAS}


def _subtipos(titulo):
    return set(re.findall(r"(?i)\btype\s*(\d|i{1,3})\b", titulo or ""))


@_registrada
def vigencia(pmid, maximo=8):
    """¿Hay una versión posterior de esta guía, consenso o revisión? Devuelve los avisos de la NLM del artículo
    (actualizado en, fe de erratas, retractado) y los candidatos posteriores con un título parecido, del más al
    menos parecido. PubMed no siempre enlaza las versiones (la Surviving Sepsis Campaign 2021 no remite a la de
    2026), por eso busca por título. Confirma cada candidato leyendo su título y su resumen; las normas anuales
    de una sociedad (p. ej., Standards of Care de la ADA) son otra serie: búscalas aparte."""
    articulo = pubmed(f"{pmid}[uid]", maximo=1)
    if not articulo:
        raise ValueError(f"PMID {pmid} no encontrado en PubMed")
    a = articulo[0]
    palabras = _palabras_titulo(a["titulo"])
    distintivas = sorted(palabras - _GENERICAS, key=len, reverse=True)[:5] or sorted(palabras, key=len,
                                                                                    reverse=True)[:4]
    desde = int(a["anio"]) + 1 if (a.get("anio") or "").isdigit() else 1900
    # Basta con todas las palabras menos una: la versión nueva puede cambiar alguna («Management of hyperglycaemia
    # in type 2 diabetes, 2022» → «Management of Type 2 Diabetes, 2026»).
    grupos = [distintivas] if len(distintivas) < 3 else [distintivas[:k] + distintivas[k + 1:]
                                                          for k in range(len(distintivas))]
    tema = " OR ".join("(" + " AND ".join(f"{p}[ti]" for p in g) + ")" for g in grupos)
    consulta = f"({tema}) AND {desde}:3000[dp] NOT ({_NO_GUIA})"
    posteriores = []
    for c in pubmed(consulta, maximo=40, orden="fecha"):
        parecido = len(palabras & _palabras_titulo(c["titulo"])) / max(1, len(palabras))
        parece_guia = bool(_TIPOS_GUIA & set(c["tipos"]) or _NUCLEO_GUIA.search(c["titulo"].rstrip(". ")))
        if _subtipos(c["titulo"]) and _subtipos(a["titulo"]) and _subtipos(c["titulo"]) != _subtipos(a["titulo"]):
            continue  # otra enfermedad: «type 1 diabetes» no actualiza una guía de «type 2 diabetes»
        if (c["pmid"] != a["pmid"] and not _COMENTARIO.search(c["titulo"])
                and (parecido >= 0.8 or (parecido >= 0.5 and parece_guia))):
            posteriores.append({k: c[k] for k in ("pmid", "titulo", "revista", "anio", "tipos", "estado", "avisos",
                                                  "pmcid", "doi")} | {"parecido": round(parecido, 2)})
    posteriores.sort(key=lambda c: (c["parecido"], c["anio"] or ""), reverse=True)
    return {"pmid": a["pmid"], "titulo": a["titulo"], "anio": a["anio"], "avisos": a["avisos"],
            "posteriores": posteriores[:maximo]}


_IRIS = {"oms": "https://iris.who.int", "ops": "https://iris.paho.org"}


def _iris_pdf(item):
    """Dirección del primer PDF de un documento de IRIS (paquete ORIGINAL de DSpace)."""
    for paquete in json.loads(_get(item["_links"]["bundles"]["href"]))["_embedded"]["bundles"]:
        if paquete.get("name") == "ORIGINAL":
            archivos = json.loads(_get(paquete["_links"]["bitstreams"]["href"]))["_embedded"]["bitstreams"]
            for archivo in archivos:
                if (archivo.get("name") or "").lower().endswith(".pdf"):
                    return archivo["_links"]["content"]["href"]
    return None


@_registrada
def iris(consulta, repositorio="oms", maximo=6, pdf=True, orden="relevancia"):
    """Documentos de la OMS (iris.who.int) o de la OPS (repositorio="ops", iris.paho.org): guías, manuales y
    documentos técnicos, con fecha, idioma, ficha y la dirección de su PDF para leerlo con pdf_texto. Busca en el
    idioma del documento («neumonía» en la OPS, «pneumonia» en la OMS). orden="fecha" pone primero lo más
    reciente, aunque sea menos pertinente: úsalo con términos precisos."""
    base = _IRIS[repositorio]
    ordenar = "&sort=dc.date.issued,DESC" if orden == "fecha" else ""
    d = json.loads(_get(f"{base}/server/api/discover/search/objects?query={urllib.parse.quote(consulta)}"
                        f"&size={maximo}&dsoType=ITEM{ordenar}"))
    salida = []
    for objeto in d["_embedded"]["searchResult"]["_embedded"]["objects"]:
        item = objeto["_embedded"]["indexableObject"]
        metadatos = item.get("metadata", {})

        def valor(clave):
            return (metadatos.get(clave) or [{}])[0].get("value")
        r = {"titulo": valor("dc.title"), "fecha": valor("dc.date.issued"), "idioma": valor("dc.language.iso"),
             "autor": valor("dc.contributor.author") or valor("dc.contributor.corpauthor"),
             "isbn": valor("dc.identifier.isbn"),
             "url": valor("dc.identifier.uri") or f"{base}/handle/{item.get('handle')}"}
        if pdf:
            try:
                r["pdf"] = _iris_pdf(item)
            except (RuntimeError, KeyError):
                r["pdf"] = None
        salida.append(r)
    return salida


def _plano(texto):
    import unicodedata
    return unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode().lower()


@_registrada
def binasss(termino, maximo=20):
    """Documentos de la Biblioteca Nacional de Salud y Seguridad Social (BINASSS, CCSS, Costa Rica): protocolos,
    normas y guías de la CCSS y documentos que la biblioteca aloja (a veces, copias de guías internacionales).
    Busca en el buscador del sitio y en su página de protocolos, normas y guías, y devuelve los títulos que
    contienen el término (sin tildes ni mayúsculas) con su enlace, casi siempre un PDF para pdf_texto."""
    buscado, salida, vistos = _plano(termino), [], set()
    for url in (f"https://www.binasss.sa.cr/?s={urllib.parse.quote(termino)}",
                "https://www.binasss.sa.cr/protocolos/protocolos.htm"):
        crudo = _get(url, cabeceras=_NAVEGADOR)
        try:
            contenido = crudo.decode("utf-8")
        except UnicodeDecodeError:
            contenido = crudo.decode("latin-1")
        for enlace, etiqueta in re.findall(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', contenido, re.S):
            titulo = html.unescape(re.sub(r"<[^>]+>|\s+", " ", etiqueta)).strip()
            direccion = urllib.parse.urljoin(url, html.unescape(enlace))
            if len(titulo) < 12 or buscado not in _plano(titulo) or direccion in vistos:
                continue
            vistos.add(direccion)
            salida.append({"titulo": titulo, "url": direccion, "pdf": direccion.lower().endswith(".pdf")})
    return salida[:maximo]


@_registrada
def pagina(url, patrones=(), contexto=300):
    """Texto de una página web pública (guía, ficha técnica nacional, página de un instituto), en cualquier idioma.
    Quita el HTML y devuelve, para cada patrón (regex en el idioma de la página), los fragmentos donde aparece.
    Si el texto es muy corto, la página probablemente se genera con JavaScript: búscala en otra fuente."""
    crudo = _get(url, cabeceras=_NAVEGADOR).decode("utf-8", "ignore")
    titulo = re.search(r"<title[^>]*>(.*?)</title>", crudo, re.S | re.I)
    idioma = re.search(r'<html[^>]*\blang="([^"]+)"', crudo, re.I)
    texto = re.sub(r"<script.*?</script>|<style.*?</style>|<noscript.*?</noscript>", " ", crudo, flags=re.S | re.I)
    texto = html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", texto))).strip()
    fragmentos = {p: [texto[max(0, m.start() - contexto):m.end() + contexto]
                      for m in list(re.finditer(p, texto, re.I))[:3]] for p in patrones}
    return {"url": url, "titulo": html.unescape(titulo.group(1).strip()) if titulo else None,
            "idioma": idioma.group(1) if idioma else None, "caracteres": len(texto),
            "aviso": "texto muy corto: puede requerir JavaScript" if len(texto) < 2000 else None,
            "fragmentos": fragmentos}


if __name__ == "__main__":
    argv = sys.argv[1:]
    if "--registro" in argv:  # anota las consultas en <carpeta>/busquedas.jsonl
        k = argv.index("--registro")
        registrar_busquedas(argv[k + 1])
        del argv[k:k + 2]
    opciones = {}
    for bandera, (clave, valor) in {"--latam": ("region", "latam"), "--titulo": ("titulo", True),
                                    "--fecha": ("orden", "fecha"), "--ops": ("repositorio", "ops")}.items():
        if bandera in argv:
            argv.remove(bandera)
            opciones[clave] = valor
    accion, *args = argv
    if accion in ("pubchem", "chembl", "pdb-buscar", "bioicons", "dailymed", "cima", "reactome", "nci", "livertox",
                  "medlineplus", "lactmed", "openfda", "openfda-eventos", "cpic", "fda-indicaciones", "mesh", "mondo",
                  "guias", "guias-titulo", "togopic", "commons", "iris", "binasss", "ema", "ensayos"):
        args = [" ".join(args)]
    funciones = {"pubchem": pubchem, "chembl": chembl, "pdb-buscar": pdb_buscar, "pdb-ligandos": pdb_ligandos,
                 "bioicons": bioicons, "servier-kits": servier_kits, "servier-diapositivas": servier_diapositivas,
                 "servier-extraer": lambda k, d, g, n: servier_extraer(k, int(d), g, n),
                 "dailymed": dailymed, "cima": cima, "uniprot": uniprot, "reactome": reactome,
                 "uniprot-proteina": lambda nombre, organismo=None: uniprot(
                     proteina=nombre, organismo=int(organismo) if organismo else None),
                 "pdf": lambda url, *patrones: pdf_texto(url, patrones),
                 "pdf-enlace": pdf_desde_pagina,
                 "europepmc": lambda consulta, patron=None: europepmc(consulta, patron),
                 "pubmed": lambda consulta, patron=None: pubmed(
                     consulta, patron, **{k: v for k, v in opciones.items() if k in ("region", "orden")}),
                 "gen": ncbi_gene, "nci": nci_tesauro, "livertox": livertox, "medlineplus": medlineplus,
                 "lactmed": lactmed, "openfda": openfda, "openfda-eventos": openfda_eventos, "cpic": cpic,
                 "actividad": lambda nombre, diana=None: chembl_actividad(nombre, diana),
                 "bindingdb": bindingdb, "epar": ema_epar, "fda-indicaciones": fda_indicaciones,
                 "ema": ema_medicamento, "ensayos": ensayos, "ensayo": ensayo,
                 "mesh": mesh, "mondo": mondo,
                 "guias": lambda enfermedad: guias(
                     enfermedad, **{k: v for k, v in opciones.items() if k in ("region", "titulo")}),
                 "guias-titulo": lambda enfermedad: guias(enfermedad, titulo=True),
                 "vigencia": vigencia,
                 "iris": lambda consulta: iris(
                     consulta, **{k: v for k, v in opciones.items() if k in ("repositorio", "orden")}),
                 "binasss": binasss,
                 "togopic": togopic, "togopic-descargar": togopic_descargar, "commons": commons,
                 "commons-descargar": commons_descargar,
                 "pmc": lambda pmcid, *patrones: pmc_texto(pmcid, patrones),
                 "texto": lambda identificador, *patrones: texto_completo(identificador, patrones),
                 "acceso-abierto": lambda doi, *patrones: acceso_abierto(doi, patrones),
                 "ids": ids,
                 "pagina": lambda url, *patrones: pagina(url, patrones)}
    resultado = funciones[accion](*args)
    print(json.dumps(resultado, indent=2, ensure_ascii=False, default=str))
